from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "presentations" / "2026-06-23_progress_shimada_redesign.pptx"
SHOT = ROOT / "Assets" / "Screenshots" / "v8_10_camera_radius_13_2_play_20s.png"
OVERVIEW = ROOT / "Assets" / "Screenshots" / "v8_10_overview_switchback_incidents.png"

W, H = Inches(13.333), Inches(7.5)

BG = RGBColor(242, 238, 229)
PANEL = RGBColor(252, 249, 242)
INK = RGBColor(35, 36, 43)
MUTED = RGBColor(91, 88, 80)
LINE = RGBColor(196, 185, 168)
ORANGE = RGBColor(238, 112, 58)
TEAL = RGBColor(47, 170, 177)
GREEN = RGBColor(111, 153, 90)
DARK = RGBColor(75, 72, 66)
RED = RGBColor(203, 77, 60)
YELLOW = RGBColor(241, 186, 88)

FONT = "Yu Gothic"


def add_text(slide, x, y, w, h, text, size=24, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    for run in p.runs:
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_label(slide, text, color=ORANGE):
    add_text(slide, 0.48, 0.10, 2.6, 0.28, text, size=10, bold=True, color=color)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.48), Inches(0.34), Inches(1.10), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()


def add_page(slide, n):
    add_text(slide, 12.78, 7.12, 0.42, 0.22, str(n), size=9, color=MUTED, align=PP_ALIGN.RIGHT)


def add_card(slide, x, y, w, h, title, body, accent=ORANGE, body_size=15):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = PANEL
    shp.line.color.rgb = LINE
    shp.line.width = Pt(1)
    add_text(slide, x + 0.22, y + 0.18, w - 0.44, 0.34, title, size=17, bold=True, color=accent)
    add_text(slide, x + 0.22, y + 0.66, w - 0.44, h - 0.82, body, size=body_size, color=INK)


def add_bullets(slide, x, y, w, h, bullets, size=17, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.space_after = Pt(7)
        p.level = 0
        for run in p.runs:
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.color.rgb = color
    return box


def add_image(slide, path, x, y, w, h):
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    else:
        add_card(slide, x, y, w, h, "画像未生成", path.name, TEAL)


def add_arrow(slide, x1, y1, x2, y2, color=MUTED):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(1.4)
    return c


def add_chip(slide, x, y, w, text, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.42))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    add_text(slide, x, y + 0.08, w, 0.22, text, size=11, bold=True, color=RGBColor(255, 255, 255), align=PP_ALIGN.CENTER)


def new_slide(prs, page=None, label=None, label_color=ORANGE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    if label:
        add_label(slide, label, label_color)
    if page is not None:
        add_page(slide, page)
    return slide


prs = Presentation()
prs.slide_width = W
prs.slide_height = H

# 1. Title
s = new_slide(prs)
add_text(s, 0.62, 0.58, 11.7, 0.72, "進捗報告", size=44, bold=True)
add_text(s, 0.66, 1.34, 10.6, 0.46, "Shrinking Island の現状と、ゲームを作り直す方針", size=25, color=MUTED)
add_image(s, SHOT, 6.45, 2.00, 6.25, 3.55)
add_text(s, 0.70, 5.78, 4.6, 0.34, "島田 泰成", size=20, bold=True)
add_text(s, 0.70, 6.20, 4.6, 0.28, "2026-06-23", size=13, color=MUTED)
add_chip(s, 0.70, 2.25, 1.45, "現状", TEAL)
add_text(s, 0.70, 2.83, 4.92, 1.00, "M2 の足回りと Shrinking Island の土台は進んだ。\nただし、遊びとしては設計を見直す必要が出てきた。", size=22, bold=True)

# 2. Current state
s = new_slide(prs, 2, "CURRENT")
add_text(s, 0.55, 0.56, 12.1, 0.55, "ここまでは、M2 の足回りを締めて M3 検証へ進める土台を作った", size=30, bold=True)
add_card(s, 0.62, 1.35, 3.75, 1.65, "プレイ体感", "足・顔・移動・カメラの違和感を実測ベースで修正。\n酔いやすい camera 演出は撤退。", ORANGE)
add_card(s, 4.78, 1.35, 3.75, 1.65, "縮む島 V8.10", "4隅スポーン、外周ルート、橋、中央圧力、復帰 launcher まで再設計。", TEAL)
add_card(s, 8.94, 1.35, 3.75, 1.65, "品質管理", "Console 0、EditMode 9/9、docs / review に直した理由を残した。", GREEN)
add_image(s, OVERVIEW, 0.72, 3.42, 6.05, 3.10)
add_bullets(s, 7.22, 3.55, 5.15, 2.58, [
    "直線で中央へ行きづらくするため、switchback と incident pocket を追加。",
    "中心へ向かうほど接触しやすい構造を狙った。",
    "発表用デッキ、録画チェックリスト、ハンドオフを更新済み。"
], size=18)

# 3. Problem
s = new_slide(prs, 3, "PROBLEM", RED)
add_text(s, 0.55, 0.56, 12.1, 0.55, "でも、今のままだとゲームとして面白くない可能性が高い", size=30, bold=True)
add_card(s, 0.65, 1.45, 3.8, 3.35, "問題 1", "広くしただけだと、移動時間が伸びるだけ。\n\nプレイヤー同士が自然にぶつかる理由が弱い。", RED, body_size=18)
add_card(s, 4.85, 1.45, 3.8, 3.35, "問題 2", "障害物を足すほど、画面はリッチになるが遊びの目的が読みにくくなる。\n\n今は情報量が多い。", RED, body_size=18)
add_card(s, 9.05, 1.45, 3.8, 3.35, "問題 3", "中央が「安全地帯」に見えると、全員が同じ行動をして終わる。\n\n駆け引きが薄い。", RED, body_size=18)
add_text(s, 0.90, 5.66, 11.50, 0.62, "判断: V8 を磨き続けるより、V9 として遊びのルールから設計し直す。", size=23, bold=True, color=INK)

# 4. Diagnosis
s = new_slide(prs, 4, "DIAGNOSIS", ORANGE)
add_text(s, 0.55, 0.56, 12.1, 0.55, "足りないのは「ステージの広さ」ではなく、事故が起きる必然性", size=30, bold=True)
add_card(s, 0.72, 1.42, 3.4, 1.35, "現状の構造", "安全な中央へ進む", TEAL, body_size=20)
add_arrow(s, 4.25, 2.10, 5.25, 2.10)
add_card(s, 5.45, 1.42, 3.4, 1.35, "起きやすい行動", "他人を避けて移動", TEAL, body_size=20)
add_arrow(s, 8.95, 2.10, 9.95, 2.10)
add_card(s, 10.10, 1.42, 2.55, 1.35, "結果", "作業感", RED, body_size=20)
add_card(s, 0.72, 3.55, 3.4, 1.35, "必要な構造", "安全に行きたい", ORANGE, body_size=20)
add_arrow(s, 4.25, 4.22, 5.25, 4.22, ORANGE)
add_card(s, 5.45, 3.55, 3.4, 1.35, "でも", "同じ通路に押し込まれる", ORANGE, body_size=20)
add_arrow(s, 8.95, 4.22, 9.95, 4.22, ORANGE)
add_card(s, 10.10, 3.55, 2.55, 1.35, "結果", "事故", GREEN, body_size=20)
add_text(s, 0.85, 5.78, 11.2, 0.45, "設計の中心文: 中央へ逃げたい。でも安全ルートほど他プレイヤーとぶつかる。", size=21, bold=True)

# 5. V9 concept
s = new_slide(prs, 5, "V9 CONCEPT", TEAL)
add_text(s, 0.55, 0.56, 12.0, 0.55, "V9 は「強制合流するバトルコース」に作り直す", size=29, bold=True)
add_card(s, 0.70, 1.34, 3.85, 3.60, "1. 強制合流", "開始は4隅。\n10秒以内に、2人以上が同じ橋・通路・中央入口へ入るようにする。", TEAL, body_size=18)
add_card(s, 4.75, 1.34, 3.85, 3.60, "2. 中央を奪い合う", "中央はただの安全地帯にしない。\n触るとゲージ、回復、スコアなどの報酬を置く。", ORANGE, body_size=18)
add_card(s, 8.80, 1.34, 3.85, 3.60, "3. スキル用途を固定", "dash は押し出し。\njump は段差・崩落越え。\npulse は中央と橋の制圧。", GREEN, body_size=18)
add_text(s, 0.82, 5.55, 11.6, 0.52, "見た目は広く、実際の接触ポイントは狭く。遊びは「逃げる」ではなく「逃げたいのにぶつかる」。", size=21, bold=True)

# 6. New round flow
s = new_slide(prs, 6, "ROUND FLOW", GREEN)
add_text(s, 0.55, 0.56, 12.1, 0.55, "新しい 1 ラウンドの流れ", size=31, bold=True)
steps = [
    ("0-5s", "4隅スタート", "目的地と危険ルートを読む"),
    ("5-15s", "外周崩落", "同じ橋へ押し出される"),
    ("15-30s", "中央争奪", "dash / pulse で押し合う"),
    ("30s-", "最終圧縮", "jump / 復帰 / KO が連鎖")
]
xs = [0.72, 3.90, 7.08, 10.26]
colors = [TEAL, ORANGE, GREEN, RED]
for i, (time, title, body) in enumerate(steps):
    add_chip(s, xs[i], 1.42, 1.15, time, colors[i])
    add_card(s, xs[i], 1.95, 2.55, 2.25, title, body, colors[i], body_size=17)
    if i < 3:
        add_arrow(s, xs[i] + 2.65, 3.05, xs[i + 1] - 0.15, 3.05, colors[i])
add_card(s, 0.72, 5.15, 11.95, 0.92, "発表での言い方", "今の V8 は土台として使うが、面白さの核は V9 で作り直す。次回は 30 秒以内に 2 回以上事故が起きる灰色ブロック版を見せる。", ORANGE, body_size=19)

# 7. Acceptance criteria
s = new_slide(prs, 7, "NEXT CHECK")
add_text(s, 0.55, 0.56, 12.1, 0.55, "次に作る V9 は、見た目より先に面白さを検証する", size=30, bold=True)
add_card(s, 0.65, 1.34, 3.85, 3.35, "合格条件", "30秒以内に接触が2回以上起きる。\n\n中央到達が作業ではなく、奪い合いになる。", GREEN, body_size=18)
add_card(s, 4.85, 1.34, 3.85, 3.35, "確認すること", "dash / jump / pulse の使いどころが、それぞれ違って見える。\n\nKO が納得できる。", TEAL, body_size=18)
add_card(s, 9.05, 1.34, 3.85, 3.35, "やらないこと", "装飾を増やして誤魔化さない。\n\n広さだけを増やさない。\n\n動画映えだけで判断しない。", RED, body_size=18)
add_text(s, 0.90, 5.58, 11.5, 0.52, "次の報告目標: V9 の灰色ブロック版で、1 ラウンドの流れと事故ポイントを動画で見せる。", size=22, bold=True)

# 8. Closing
s = new_slide(prs, 8)
add_text(s, 0.74, 0.95, 11.2, 0.72, "ありがとうございました", size=44, bold=True)
add_text(s, 0.78, 1.88, 10.8, 0.40, "次回: Shrinking Island V9 の作り直し版を、2P事故が起きる形で検証", size=22, color=MUTED)
add_image(s, SHOT, 0.78, 3.00, 5.80, 3.16)
add_card(s, 7.10, 3.00, 5.45, 3.16, "今日の結論", "M2 の足回りと技術土台は進んだ。\n\nただしゲームとしての面白さは、V8 の延長では弱い。\n\nV9 では「強制合流」「中央争奪」「スキル用途」を先に作る。", ORANGE, body_size=20)

prs.save(OUT)
print(OUT)
