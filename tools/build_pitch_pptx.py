"""
2026-05-19 進捗報告パワポ ビルダー v2
方針:
  - 専門用語を平易化 (Unity を知らない人にも分かる)
  - ゲーム説明 (どんなゲームか) を冒頭に必ず置く
  - 「カオスドラフトパーティー」日本語表記で統一
  - 余白を徹底的に詰める、フォント大きく
  - 各 slide に視覚要素必ず 1 つ
"""
import sys
import io
import os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

REPO = r'C:\Users\taise\Documents\GitHub\ChaosDraftParty'
OUT  = os.path.join(REPO, 'docs', 'presentations', '2026-05-19_進捗報告.pptx')

VIDEO_DIR = os.path.join(REPO, 'docs', 'pitch', 'videos')
VIDEO_1   = os.path.join(VIDEO_DIR, '01_crouch_stand.mp4')
VIDEO_2   = os.path.join(VIDEO_DIR, '02_movement.mp4')
VIDEO_3   = os.path.join(VIDEO_DIR, '03_camera_ceiling_block.mp4')

IMG_FRONT   = os.path.join(REPO, 'Assets', '_Project', 'eggybot_view_front.png')
IMG_FRONT34 = os.path.join(REPO, 'Assets', '_Project', 'eggybot_view_front34.png')
IMG_SIDE    = os.path.join(REPO, 'Assets', '_Project', 'eggybot_view_side_right.png')
IMG_TOP     = os.path.join(REPO, 'Assets', '_Project', 'eggybot_view_top.png')

# ---- 色 ----
C_TEXT    = RGBColor(0x1F, 0x1F, 0x2E)
C_SUBTEXT = RGBColor(0x55, 0x55, 0x66)
C_ACCENT  = RGBColor(0xFF, 0x5C, 0x8D)  # 鮮やかピンク
C_ACCENT2 = RGBColor(0xFF, 0xC2, 0x3C)  # 黄 (カオス感)
C_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
C_GREY    = RGBColor(0x88, 0x88, 0x99)
C_BG_DARK = RGBColor(0x22, 0x22, 0x36)
C_BG_PALE = RGBColor(0xFF, 0xEE, 0xF3)  # 薄ピンク
C_BG_BLUE = RGBColor(0xEC, 0xF3, 0xFF)  # 薄水
C_BG_YEL  = RGBColor(0xFF, 0xF8, 0xE0)  # 薄黄
C_GOOD    = RGBColor(0x3C, 0xB0, 0x4B)
C_BAD     = RGBColor(0xE5, 0x39, 0x35)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW = prs.slide_width
SH = prs.slide_height


# ---------- helpers ----------
def rect(slide, x, y, w, h, color, line=False, line_color=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    if line:
        sh.line.color.rgb = line_color or C_TEXT
        sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh

def rrect(slide, x, y, w, h, color, radius=0.25):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh

def text(slide, x, y, w, h, content, size=18, bold=False, color=C_TEXT,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font='Yu Gothic UI',
         line_spacing=1.15):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left   = Inches(0.04)
    tf.margin_right  = Inches(0.04)
    tf.margin_top    = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    lines = content.split('\n') if isinstance(content, str) else content
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return tb

def header(slide, title_str):
    rect(slide, 0, 0, SW, Inches(0.78), C_BG_DARK)
    text(slide, Inches(0.4), Inches(0.08), Inches(12.5), Inches(0.62),
         title_str, size=28, bold=True, color=C_WHITE,
         anchor=MSO_ANCHOR.MIDDLE)
    rect(slide, 0, Inches(0.78), SW, Inches(0.06), C_ACCENT)

def page_num(slide, n, total):
    text(slide, Inches(12.4), Inches(7.18), Inches(0.85), Inches(0.3),
         f'{n} / {total}', size=11, color=C_GREY, align=PP_ALIGN.RIGHT)

def bullets(slide, x, y, w, h, items, size=17, color=C_TEXT, bullet='● ',
            line_spacing=1.25, space_after=4):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        # bullet を accent 色で出す
        r1 = p.add_run()
        r1.text = bullet
        r1.font.size = Pt(size)
        r1.font.bold = True
        r1.font.color.rgb = C_ACCENT
        r1.font.name = 'Yu Gothic UI'
        # 本文
        r2 = p.add_run()
        r2.text = it
        r2.font.size = Pt(size)
        r2.font.color.rgb = color
        r2.font.name = 'Yu Gothic UI'
    return tb

def add_movie(slide, x, y, w, h, path, label_under=None):
    if not os.path.exists(path):
        rect(slide, x, y, w, h, C_GREY)
        text(slide, x, y, w, h, f'[動画見つかりません]\n{os.path.basename(path)}',
             size=14, color=C_WHITE, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
        return None
    mv = slide.shapes.add_movie(path, x, y, w, h, mime_type='video/mp4')
    if label_under:
        text(slide, x, y + h + Inches(0.05), w, Inches(0.3),
             label_under, size=11, color=C_GREY, align=PP_ALIGN.CENTER)
    return mv

TOTAL = 9


# =========================================================
# Slide 1: タイトル — 「カオスドラフトパーティーとは」を 1 画面で
# =========================================================
s = prs.slides.add_slide(BLANK)

# 大きな左カラム (タイトル + ゲーム説明)
rect(s, 0, 0, Inches(8.8), SH, C_BG_DARK)
rect(s, 0, 0, Inches(0.35), SH, C_ACCENT)  # 左 accent

# タイトル
text(s, Inches(0.7), Inches(0.55), Inches(8.0), Inches(0.6),
     '卒業制作 進捗報告', size=22, color=C_ACCENT2, bold=True)
text(s, Inches(0.7), Inches(1.15), Inches(8.0), Inches(1.6),
     'カオスドラフトパーティー', size=58, bold=True, color=C_WHITE)
text(s, Inches(0.7), Inches(2.75), Inches(8.0), Inches(0.6),
     '— 2 〜 4 人で遊ぶ オンライン 3D パーティゲーム',
     size=22, color=C_WHITE)

# ゲーム説明 (3 行で)
rrect(s, Inches(0.7), Inches(3.6), Inches(7.8), Inches(0.85), C_ACCENT, radius=0.3)
text(s, Inches(0.85), Inches(3.6), Inches(7.5), Inches(0.85),
     '1 試合 8 分 / 4 ラウンド × 90 秒 / ラウンド毎にスキルを「ドラフト」して育てる',
     size=18, bold=True, color=C_WHITE, anchor=MSO_ANCHOR.MIDDLE)

# 副題: 何を見せるか
text(s, Inches(0.7), Inches(4.65), Inches(8.0), Inches(0.6),
     '今日のテーマ', size=14, color=C_ACCENT2, bold=True)
text(s, Inches(0.7), Inches(5.05), Inches(8.0), Inches(1.6),
     '卵キャラが、オンラインで、\n自然に動くようになった話',
     size=34, bold=True, color=C_WHITE, line_spacing=1.1)

# 日付
text(s, Inches(0.7), Inches(6.75), Inches(8.0), Inches(0.5),
     '2026 年 5 月 19 日   /   島田   /   小坂研究室   /   M1 完了 + M2 中盤',
     size=14, color=C_GREY)

# 右カラム = 卵キャラ画像 (大きく)
if os.path.exists(IMG_FRONT):
    s.shapes.add_picture(IMG_FRONT, Inches(9.0), Inches(1.0),
                         width=Inches(4.2), height=Inches(4.2))
text(s, Inches(9.0), Inches(5.35), Inches(4.2), Inches(0.5),
     '「EGGYBOT (エッグボット)」', size=18, bold=True, color=C_TEXT,
     align=PP_ALIGN.CENTER)
text(s, Inches(9.0), Inches(5.85), Inches(4.2), Inches(0.4),
     '本作の操作キャラ (4 色展開)', size=13, color=C_SUBTEXT,
     align=PP_ALIGN.CENTER)
# キャラ下に色チップ 4 個
chip_y = Inches(6.4)
for i, col in enumerate([RGBColor(0xFF,0x5C,0x8D),
                          RGBColor(0x4A,0xB8,0xFF),
                          RGBColor(0x6E,0xE0,0x5A),
                          RGBColor(0xFF,0xC2,0x3C)]):
    rrect(s, Inches(9.2 + i*0.95), chip_y, Inches(0.8), Inches(0.45), col, radius=0.5)

page_num(s, 1, TOTAL)


# =========================================================
# Slide 2: どんなゲーム? (もう少し詳しく + 自分の役割)
# =========================================================
s = prs.slides.add_slide(BLANK)
header(s, 'どんなゲーム?  ―  「見た目カオスでも勝つのは上手いやつ」')

# ゲーム概要 4 マス (アイコン + 大文字)
def info_card(x, y, w, h, big, lbl, sub, bg):
    rrect(s, x, y, w, h, bg, radius=0.1)
    text(s, x, y + Inches(0.15), w, Inches(1.1),
         big, size=48, bold=True, color=C_ACCENT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x, y + Inches(1.3), w, Inches(0.45),
         lbl, size=20, bold=True, color=C_TEXT, align=PP_ALIGN.CENTER)
    text(s, x, y + Inches(1.78), w, Inches(0.6),
         sub, size=13, color=C_SUBTEXT, align=PP_ALIGN.CENTER, line_spacing=1.2)

cw = Inches(3.0); ch = Inches(2.5)
cy = Inches(1.1)
info_card(Inches(0.4),  cy, cw, ch, '2〜4',  '人数',     '友達とオンライン対戦\n(PC・Steam 配信予定)',  C_BG_PALE)
info_card(Inches(3.55), cy, cw, ch, '8 分',  '1 試合',   '90 秒 × 4 ラウンド\n+ ドラフト + リザルト',   C_BG_BLUE)
info_card(Inches(6.70), cy, cw, ch, '4 種',  'ミニゲーム', '縮む島 / 爆弾鬼 /\nしっぽ狩り (バトル/逃走)', C_BG_YEL)
info_card(Inches(9.85), cy, cw, ch, '10 種', 'スキル',   'ダッシュ / 爆弾 / 盾 など\nドラフトで選んで重ねがけ', C_BG_PALE)

# 「ドラフトって何?」帯
rrect(s, Inches(0.4), Inches(3.8), Inches(12.55), Inches(1.4), C_BG_DARK, radius=0.05)
text(s, Inches(0.7), Inches(3.9), Inches(3.2), Inches(0.5),
     '「ドラフト」とは?', size=18, bold=True, color=C_ACCENT2)
text(s, Inches(0.7), Inches(4.35), Inches(12.0), Inches(1.0),
     'ラウンドが終わるたびに「次に使うスキル」を 5 秒で 1 つ選ぶ。\n'
     '同じスキルを重ねると効果アップ。組み合わせで「コンボ」も発動 → 毎試合違うビルドが組める。',
     size=16, color=C_WHITE, line_spacing=1.35)

# 開発体制 / 期間
rrect(s, Inches(0.4), Inches(5.4), Inches(8.3), Inches(1.7), C_BG_PALE, radius=0.05)
text(s, Inches(0.7), Inches(5.5), Inches(8.0), Inches(0.5),
     '開発体制', size=18, bold=True, color=C_ACCENT)
text(s, Inches(0.7), Inches(5.95), Inches(8.0), Inches(1.1),
     '1 人開発  /  2026 年 4 月 〜 9 月  /  Unity というゲームエンジン使用\n'
     '実装サポート = Claude (AI アシスタント) と二人三脚で進行',
     size=15, color=C_TEXT, line_spacing=1.4)

# 今のフェーズ
rrect(s, Inches(8.85), Inches(5.4), Inches(4.1), Inches(1.7), C_ACCENT, radius=0.05)
text(s, Inches(9.0), Inches(5.5), Inches(3.9), Inches(0.5),
     '今ここ', size=14, bold=True, color=C_ACCENT2)
text(s, Inches(9.0), Inches(5.95), Inches(3.9), Inches(0.5),
     'M2 中盤 (5 月)', size=24, bold=True, color=C_WHITE)
text(s, Inches(9.0), Inches(6.5), Inches(3.9), Inches(0.5),
     'M1 (4月: 基盤) 完了 → M2 で操作とカメラを磨く段階',
     size=11, color=C_WHITE, line_spacing=1.3)

page_num(s, 2, TOTAL)


# =========================================================
# Slide 3: 今月できたこと(一言で) + 全体像
# =========================================================
s = prs.slides.add_slide(BLANK)
header(s, '今月できたこと  ―  ひとことで言うと')

# 大見出し
rrect(s, Inches(0.4), Inches(1.05), Inches(12.55), Inches(1.4), C_ACCENT, radius=0.05)
text(s, Inches(0.4), Inches(1.05), Inches(12.55), Inches(1.4),
     '卵キャラがオンラインで、自然に、可愛く動くようになった',
     size=36, bold=True, color=C_WHITE,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# 3 つの柱
def pillar(x, w, num, head, body, color_bg):
    rrect(s, x, Inches(2.7), w, Inches(3.7), color_bg, radius=0.08)
    # 番号バッジ
    rrect(s, x + Inches(0.3), Inches(2.9), Inches(0.85), Inches(0.85), C_ACCENT, radius=0.5)
    text(s, x + Inches(0.3), Inches(2.9), Inches(0.85), Inches(0.85),
         num, size=32, bold=True, color=C_WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + Inches(1.3), Inches(2.95), w - Inches(1.5), Inches(0.75),
         head, size=21, bold=True, color=C_TEXT, line_spacing=1.1)
    text(s, x + Inches(0.3), Inches(3.95), w - Inches(0.55), Inches(2.4),
         body, size=15, color=C_TEXT, line_spacing=1.45)

pillar(Inches(0.4), Inches(4.1), '①',
       'オンラインで\n繋がるようになった',
       '友達と同じ部屋に入って、互いのキャラが画面に表示されて動く。\n\n'
       '「縮む島」というミニゲームの試作版が動き、床が消えて落ちると同期して落下する。',
       C_BG_PALE)

pillar(Inches(4.65), Inches(4.1), '②',
       'キャラを\nちゃんと動かせる',
       'キーボード・マウスでの操作を一通り作った: \n移動 / ジャンプ / ダッシュ / しゃがみ。\n\n'
       'カプセル形のキャラがしゃがむと、蓋が開いて頭がぴょこっと出る可愛い動き付き。',
       C_BG_BLUE)

pillar(Inches(8.9), Inches(4.1), '③',
       'カメラに\n変な動きが無くなった',
       'カメラが床や壁に潜らない。狭い所でしゃがんで天井があると立てない。\n\n'
       'キャラが壁の裏に隠れたら半透明で透ける(原神とかと同じ仕組み)。',
       C_BG_YEL)

# 下帯: 次のページから動画
rrect(s, Inches(0.4), Inches(6.55), Inches(12.55), Inches(0.7), C_BG_DARK, radius=0.05)
text(s, Inches(0.4), Inches(6.55), Inches(12.55), Inches(0.7),
     '↓ 次のページから動画で実物を見せます (3 本)',
     size=18, bold=True, color=C_ACCENT2,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

page_num(s, 3, TOTAL)


# =========================================================
# Slide 4: 動画 ① キャラを動かす
# =========================================================
s = prs.slides.add_slide(BLANK)
header(s, 'デモ ①  ―  キーボード・マウスでキャラを動かす')

# 動画 (大きめ)
add_movie(s, Inches(0.4), Inches(1.05), Inches(8.2), Inches(5.7), VIDEO_2,
          label_under='移動・ジャンプ・ダッシュの 30 秒デモ')

# 右カラム — 平易な解説
text(s, Inches(8.85), Inches(1.1), Inches(4.2), Inches(0.5),
     '何ができる?', size=22, bold=True, color=C_ACCENT)

ops = [
    ('WASD キー',   '前後左右に走る'),
    ('マウス',       'カメラを回す'),
    ('Space',       'ジャンプ'),
    ('Ctrl',        'ダッシュ (速くなる)'),
    ('Shift',       'しゃがむ'),
]
y_op = 1.65
for k, v in ops:
    rrect(s, Inches(8.85), Inches(y_op), Inches(1.5), Inches(0.5),
          C_ACCENT, radius=0.3)
    text(s, Inches(8.85), Inches(y_op), Inches(1.5), Inches(0.5),
         k, size=14, bold=True, color=C_WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(10.5), Inches(y_op), Inches(2.6), Inches(0.5),
         v, size=16, color=C_TEXT, anchor=MSO_ANCHOR.MIDDLE)
    y_op += 0.65

# 工夫したこと
rrect(s, Inches(8.85), Inches(5.15), Inches(4.2), Inches(1.95), C_BG_PALE, radius=0.05)
text(s, Inches(9.0), Inches(5.2), Inches(4.0), Inches(0.5),
     '裏で工夫したこと', size=15, bold=True, color=C_ACCENT)
text(s, Inches(9.0), Inches(5.6), Inches(4.0), Inches(1.5),
     '走っても卵がフラフラしない\n'
     '(腕や体のブレを抑える仕組みを自作)\n\n'
     'カメラの向きと連動して進む方向が変わる',
     size=12, color=C_TEXT, line_spacing=1.4)

page_num(s, 4, TOTAL)


# =========================================================
# Slide 5: 動画 ② しゃがみ・立ち上がり
# =========================================================
s = prs.slides.add_slide(BLANK)
header(s, 'デモ ②  ―  しゃがみ & 立ち上がり (蓋の開閉アニメ)')

add_movie(s, Inches(0.4), Inches(1.05), Inches(8.2), Inches(5.7), VIDEO_1,
          label_under='Shift キーでしゃがむ / 離すと立ち上がる')

text(s, Inches(8.85), Inches(1.1), Inches(4.2), Inches(0.5),
     'カプセル開閉ギミック', size=22, bold=True, color=C_ACCENT)
text(s, Inches(8.85), Inches(1.7), Inches(4.2), Inches(1.4),
     'カプセル型 (ガチャポン) のキャラ\nなので「蓋を開けて頭が出る」\nアニメを自作。',
     size=15, color=C_TEXT, line_spacing=1.4)

# 4 step 図解
rrect(s, Inches(8.85), Inches(3.2), Inches(4.2), Inches(3.5), C_BG_BLUE, radius=0.05)
text(s, Inches(9.0), Inches(3.25), Inches(4.0), Inches(0.5),
     'しゃがみの動き 4 段階', size=15, bold=True, color=C_ACCENT)
steps = [
    ('1', 'ぴょん!  頭が上に飛ぶ + 蓋オープン'),
    ('2', '頂点でちょっと止まる (タメ)'),
    ('3', 'スポン!  頭が落ちる + 蓋同時に閉まる'),
    ('4', '蓋がバウンドして反動 (-8° → +5° → 0°)'),
]
ys = 3.7
for n, d in steps:
    rrect(s, Inches(9.0), Inches(ys), Inches(0.5), Inches(0.5), C_ACCENT, radius=0.5)
    text(s, Inches(9.0), Inches(ys), Inches(0.5), Inches(0.5),
         n, size=18, bold=True, color=C_WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(9.65), Inches(ys), Inches(3.35), Inches(0.5),
         d, size=12, color=C_TEXT, anchor=MSO_ANCHOR.MIDDLE)
    ys += 0.7

page_num(s, 5, TOTAL)


# =========================================================
# Slide 6: 動画 ③ カメラの工夫
# =========================================================
s = prs.slides.add_slide(BLANK)
header(s, 'デモ ③  ―  カメラの工夫 (床に潜らない・天井で頭ぶつける)')

add_movie(s, Inches(0.4), Inches(1.05), Inches(8.2), Inches(5.7), VIDEO_3,
          label_under='カメラを下に向けると床手前で止まる / 天井下でしゃがむと立てない')

text(s, Inches(8.85), Inches(1.1), Inches(4.2), Inches(0.5),
     '3 つの仕組み', size=22, bold=True, color=C_ACCENT)

cam_items = [
    ('A', '床に潜らない',
     'カメラを下に向けすぎても、床に当たる手前で\n自動的に止まる (常時計算)'),
    ('B', '壁裏で半透明',
     '壁の向こうにキャラが行ったら、ドット柄で\n透けて見える (原神・モンハンと同じ仕組み)'),
    ('C', '天井で頭ぶつけ',
     '低い天井の下でしゃがんで、立ち上がろうと\nしたら頭が当たって立てない (反動付き)'),
]
yy = 1.7
for tag, ttl, body in cam_items:
    rrect(s, Inches(8.85), Inches(yy), Inches(4.2), Inches(1.65), C_BG_PALE, radius=0.05)
    rrect(s, Inches(9.0), Inches(yy + 0.1), Inches(0.55), Inches(0.55), C_ACCENT, radius=0.4)
    text(s, Inches(9.0), Inches(yy + 0.1), Inches(0.55), Inches(0.55),
         tag, size=22, bold=True, color=C_WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(9.7), Inches(yy + 0.1), Inches(3.3), Inches(0.55),
         ttl, size=15, bold=True, color=C_TEXT, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(9.0), Inches(yy + 0.75), Inches(4.0), Inches(0.85),
         body, size=11, color=C_TEXT, line_spacing=1.3)
    yy += 1.78

page_num(s, 6, TOTAL)


# =========================================================
# Slide 7: 良かったこと 5 つ (平易表現)
# =========================================================
s = prs.slides.add_slide(BLANK)
header(s, '良かったこと  ―  5 月「やってよかった」5 つ')

good = [
    ('オンライン対戦の基盤が出来た',
     'ネット越しに 2 人が同じ画面で遊べるところまで完成。\n業界標準のやり方 (まずローカル → 後でネット越し) で着実に進めた。'),
    ('カメラの質を AAA タイトル並みに前倒し',
     '本来 6 月予定だった「キャラが壁裏で透ける」仕組みを 5 月に実装。\n原神・モンハン・FF14 と同じやり方 (ドット柄での半透明) を採用。'),
    ('「失敗しないための知識ベース」を作った',
     '今までの試行錯誤を文書化 (レシピ / 落とし穴 / コード雛形)。\n次のプロジェクトでも同じ手順で AI に作業させられるようになった。'),
    ('AI 同士でコードレビューする仕組み',
     'コードを書いたら、別の AI に「ここダメじゃない?」とチェックさせる。\n人の目を通さずに、commit 前に品質を上げる流れを定着。'),
    ('「気持ちいい動き」の作り方を一般化',
     '一瞬パッと反応 + その後ぬるっと、という普遍ルールを発見。\nダメージ表示・カメラ揺れ・ボタン押し全てに応用可能。'),
]

base_y = 1.05
row_h = Inches(1.18)
for i, (ttl, desc) in enumerate(good):
    y = Inches(1.05 + i * 1.22)
    rect(s, Inches(0.4), y, Inches(12.55), row_h, C_BG_PALE)
    # 番号バッジ (大きく)
    rect(s, Inches(0.4), y, Inches(1.1), row_h, C_GOOD)
    text(s, Inches(0.4), y, Inches(1.1), row_h,
         f'0{i+1}', size=42, bold=True, color=C_WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(1.7), y + Inches(0.08), Inches(11.0), Inches(0.45),
         ttl, size=20, bold=True, color=C_TEXT)
    text(s, Inches(1.7), y + Inches(0.55), Inches(11.0), Inches(0.65),
         desc, size=14, color=C_SUBTEXT, line_spacing=1.35)

page_num(s, 7, TOTAL)


# =========================================================
# Slide 8: 失敗 5 つ (平易表現)
# =========================================================
s = prs.slides.add_slide(BLANK)
header(s, '反省  ―  うまくいかなかった 5 つ')

bad = [
    ('キャラの黒い縁取り、2 回失敗',
     'アニメ調にする「黒の輪郭線」を 2 種類の方法で試したが両方失敗。\n1 回目: 線の濃さ調整が AI 単独で詰められず。2 回目: 拾った OSS が Unity 6 で動かず。明日 3 度目の挑戦予定。'),
    ('Unity 6.4 (最新版) の仕様変更に翻弄',
     '最新 Unity で、これまで動いてた拡張機能が動かないケース複数。\n古い書き方をしてる OSS は軒並み Unity 6 で死んでた。罠を文書化して次回回避。'),
    ('効果音 / BGM が後回しに',
     'カメラ酔い対策を優先したので、音周りが手付かず。\n素材リスト・接続手順は記録済、来週着手予定。'),
    ('開発ルール (並列調査 / AI レビュー) を 2 回端折った',
     '急いで実装に走った時に「ちゃんと調べてから」と本人 (User) に 2 回指摘された。\nルール忘れを防ぐため、作業前のチェック step を作業フローに組み込み中。'),
    ('カメラ感度の調整に時間溶かした',
     '「マウスをゆっくり動かすと感度が高い、早く動かすと遅い」現象。\n調整を盛り過ぎてた → 全部外したらキレイになった。「シンプル is best」を再確認。'),
]

for i, (ttl, desc) in enumerate(bad):
    y = Inches(1.05 + i * 1.22)
    rect(s, Inches(0.4), y, Inches(12.55), Inches(1.18),
         RGBColor(0xFF, 0xEC, 0xEC))
    rect(s, Inches(0.4), y, Inches(1.1), Inches(1.18), C_BAD)
    text(s, Inches(0.4), y, Inches(1.1), Inches(1.18),
         f'0{i+1}', size=42, bold=True, color=C_WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(1.7), y + Inches(0.08), Inches(11.0), Inches(0.45),
         ttl, size=20, bold=True, color=C_TEXT)
    text(s, Inches(1.7), y + Inches(0.55), Inches(11.0), Inches(0.65),
         desc, size=14, color=C_SUBTEXT, line_spacing=1.35)

page_num(s, 8, TOTAL)


# =========================================================
# Slide 9: 数字 + 次のステップ + 締め (1 枚に集約)
# =========================================================
s = prs.slides.add_slide(BLANK)
header(s, '数字で見る 5 月  &  次のステップ')

# 数字エリア (左半分)
text(s, Inches(0.4), Inches(1.0), Inches(6.2), Inches(0.5),
     '5 月の活動を数字で', size=18, bold=True, color=C_ACCENT)

stats = [
    ('13', 'コミット数'),
    ('8',  '試行回数 (カメラ半透明)'),
    ('5',  '日報の本数'),
    ('3',  '今日撮った動画'),
    ('2',  '失敗 (黒輪郭)'),
    ('0',  '目標未達'),
]
sw = Inches(2.0); sh_ = Inches(1.45)
for i, (big, lbl) in enumerate(stats):
    r = i // 3; c = i % 3
    x = Inches(0.4 + c * 2.1)
    y = Inches(1.55 + r * 1.55)
    rrect(s, x, y, sw, sh_, C_BG_PALE, radius=0.08)
    text(s, x, y + Inches(0.05), sw, Inches(0.85),
         big, size=44, bold=True, color=C_ACCENT,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x, y + Inches(0.9), sw, Inches(0.45),
         lbl, size=12, color=C_TEXT, bold=True, align=PP_ALIGN.CENTER)

# 次のステップ (右半分)
rrect(s, Inches(6.9), Inches(1.0), Inches(6.05), Inches(3.6), C_BG_BLUE, radius=0.08)
text(s, Inches(7.1), Inches(1.1), Inches(5.7), Inches(0.5),
     '今週〜来週やること', size=18, bold=True, color=C_ACCENT)
bullets(s, Inches(7.1), Inches(1.65), Inches(5.7), Inches(2.9), [
    'キャラの黒い輪郭 (3 度目の挑戦)',
    '走る時の砂埃・着地の輪・足跡',
    'ダッシュ時の演出 (画面が引き伸び)',
    '効果音 / BGM の組み込み',
    'キャラが押される時の伸び縮み',
], size=15, line_spacing=1.4, space_after=2)

rrect(s, Inches(6.9), Inches(4.7), Inches(6.05), Inches(2.4), C_BG_YEL, radius=0.08)
text(s, Inches(7.1), Inches(4.8), Inches(5.7), Inches(0.5),
     '来月 (6 月) からやること', size=18, bold=True, color=C_ACCENT)
bullets(s, Inches(7.1), Inches(5.35), Inches(5.7), Inches(1.7), [
    'スキル 3 種実装 (ダッシュ / 弾 / 盾)',
    'ドラフト画面 (5 秒で選ぶ)',
    '4 ラウンド構造の確立',
    'ミニゲーム +2 種追加',
], size=15, line_spacing=1.4, space_after=2)

# 締め band
rect(s, 0, Inches(7.15), SW, Inches(0.35), C_BG_DARK)
text(s, Inches(0.4), Inches(7.15), Inches(12.55), Inches(0.35),
     '9 月の完成 (Steam 配信 + 外部テスト) に向けて順調に進行中',
     size=13, bold=True, color=C_ACCENT2,
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

page_num(s, 9, TOTAL)


# =========================================================
prs.save(OUT)
print(f'[OK] Saved: {OUT}')
print(f'Slides: {len(prs.slides)}')
for v in [VIDEO_1, VIDEO_2, VIDEO_3]:
    print(f'  {"✓" if os.path.exists(v) else "✗"} {os.path.basename(v)}')
