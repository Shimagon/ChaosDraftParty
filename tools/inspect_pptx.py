"""Inspect existing pptx for template reuse."""
import sys
import io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pptx import Presentation
from pptx.util import Emu

path = r'C:\Users\taise\Documents\GitHub\ChaosDraftParty\docs\presentations\2026-04-28_キャラクター総括.pptx'
p = Presentation(path)
print(f'Slide size: {p.slide_width/914400:.2f} x {p.slide_height/914400:.2f} inches')
print(f'Slides: {len(p.slides)}')
print('=' * 60)
for i, s in enumerate(p.slides):
    print(f'\n## Slide {i+1} (layout: {s.slide_layout.name})')
    for shape in s.shapes:
        kind = shape.shape_type
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                runs = []
                for run in para.runs:
                    fs = run.font.size.pt if run.font.size else '?'
                    bold = 'B' if run.font.bold else ''
                    color = ''
                    try:
                        if run.font.color and run.font.color.rgb:
                            color = f'#{run.font.color.rgb}'
                    except Exception:
                        pass
                    runs.append(f'[{fs}pt{bold}{color}]{run.text}')
                if runs:
                    print('  TEXT:', ' '.join(runs))
        elif kind == 13:
            print(f'  PIC: {shape.name} ({shape.width/914400:.2f}x{shape.height/914400:.2f})')
        else:
            print(f'  SHAPE({kind}): {shape.name}')
