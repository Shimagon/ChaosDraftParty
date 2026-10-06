from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "presentations" / "2026-06-09_progress_shimada_v8_10_autonomous.pptx"
SHOT = ROOT / "Assets" / "Screenshots" / "v8_10_camera_radius_13_2_play_20s.png"
OVERVIEW = ROOT / "Assets" / "Screenshots" / "v8_10_overview_switchback_incidents.png"

W, H = Inches(13.333), Inches(7.5)
BG = RGBColor(242, 238, 229)
INK = RGBColor(36, 37, 44)
MUTED = RGBColor(92, 88, 80)
ORANGE = RGBColor(238, 112, 58)
TEAL = RGBColor(51, 172, 178)
GREEN = RGBColor(112, 154, 92)
DARK = RGBColor(77, 73, 66)
CARD = RGBColor(252, 249, 242)
LINE = RGBColor(198, 187, 170)


def add_text(slide, x, y, w, h, text, size=24, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    for run in p.runs:
        run.font.name = "Yu Gothic"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_label(slide, text, color=ORANGE):
    add_text(slide, 0.48, 0.10, 2.2, 0.28, text, size=10, bold=True, color=color)
    shape = slide.shapes.add_shape(1, Inches(0.48), Inches(0.34), Inches(1.10), Inches(0.03))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.color.rgb = color


def add_page(slide, n):
    add_text(slide, 12.85, 7.18, 0.35, 0.18, str(n), size=9, color=MUTED, align=PP_ALIGN.RIGHT)


def add_card(slide, x, y, w, h, title, body, accent=ORANGE):
    shp = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = CARD
    shp.line.color.rgb = LINE
    shp.shadow.inherit = False
    add_text(slide, x + 0.22, y + 0.18, w - 0.44, 0.30, title, size=18, bold=True, color=accent)
    add_text(slide, x + 0.22, y + 0.65, w - 0.44, h - 0.80, body, size=14, color=INK)


def add_bullets(slide, x, y, w, h, bullets, size=17):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.level = 0
        p.space_after = Pt(8)
        for run in p.runs:
            run.font.name = "Yu Gothic"
            run.font.size = Pt(size)
            run.font.color.rgb = INK
    return box


def add_image(slide, path, x, y, w, h):
    if path.exists():
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    else:
        add_card(slide, x, y, w, h, "画像未生成", str(path), accent=TEAL)


prs = Presentation()
prs.slide_width = W
prs.slide_height = H
blank = prs.slide_layouts[6]


def new_slide(page=None, label=None):
    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    if label:
        add_label(s, label)
    if page is not None:
        add_page(s, page)
    return s


# 1 title
s = new_slide()
add_text(s, 0.65, 0.70, 11.8, 0.55, "進捗報告", size=42, bold=True)
add_text(s, 0.70, 1.45, 8.8, 0.45, "Shrinking Island の足回りとステージ再設計", size=25, color=MUTED)
add_text(s, 0.72, 6.10, 4.2, 0.32, "島田 泰成", size=20, bold=True)
add_text(s, 0.72, 6.55, 3.5, 0.24, "2026-06-09", size=13, color=MUTED)
add_image(s, SHOT, 6.25, 2.00, 6.45, 3.65)


# 2 context
s = new_slide(2, "CONTEXT")
add_text(s, 0.55, 0.55, 12.2, 0.55, "Shrinking Island は「中央へ逃げるほど、ぶつかる」ミニゲーム", size=28, bold=True)
add_card(s, 0.65, 1.45, 3.8, 2.0, "何をするゲームか", "外側の足場が順番に消える。\nプレイヤーは中央の安全地帯へ移動する。", TEAL)
add_card(s, 4.75, 1.45, 3.8, 2.0, "面白さの狙い", "安全な道ほど狭くなる。\n橋・障害物・中央入口で接触が起きる。", ORANGE)
add_card(s, 8.85, 1.45, 3.8, 2.0, "今回の問題", "前の版は、初見だと目的が読みにくく、すぐ中央へ行けた。", GREEN)
add_bullets(s, 1.05, 4.45, 11.3, 1.20, [
    "発表では最初に「足場が消える中、中央へ逃げるゲーム」と説明する。",
    "ただし、ただ逃げるだけではなく、他プレイヤーと事故が起きる構造にした。"
], size=18)


# 3 today
s = new_slide(3, "TODAY")
add_text(s, 0.55, 0.55, 12.2, 0.55, "M2 の足回りを締め、M3 の遊び検証へ進める状態にした", size=28, bold=True)
add_card(s, 0.55, 1.35, 5.9, 1.55, "プレイ体感", "足・顔・移動・カメラの違和感を実測ベースで修正。\n酔いやすい camera 演出は撤退。", ORANGE)
add_card(s, 6.85, 1.35, 5.9, 1.55, "縮む島", "四隅スタート、橋、中央圧力、復帰 launcher の構造へ再設計。", ORANGE)
add_card(s, 0.55, 3.25, 5.9, 1.55, "見た目とVFX", "砂埃、警告、fade、launcher、sky/post を Unity 臭が出にくい方向へ更新。", ORANGE)
add_card(s, 6.85, 3.25, 5.9, 1.55, "品質管理", "AI review loop と Codex handoff を導入。\n誤診の原因を docs / review に残した。", ORANGE)


# 4 stage
s = new_slide(4, "STAGE")
add_text(s, 0.55, 0.52, 12.1, 0.55, "直線で中央へ行くだけにならないよう、外周ルートを作った", size=28, bold=True)
add_image(s, OVERVIEW, 0.62, 1.24, 7.0, 4.55)
add_bullets(s, 8.00, 1.45, 4.7, 3.7, [
    "四隅に安定したスポーンデッキを追加",
    "外周にジグザグ障害物と目印を追加",
    "高台を左右に分割し、スキルの着地点を残した",
    "中央の黒い塊を軽くして、構造として読ませた"
], size=17)
add_card(s, 8.00, 5.25, 4.7, 0.95, "次に見ること", "2P で中盤ポケットと中央入口が事故を増やすか。", TEAL)


# 5 spec
s = new_slide(5, "SPEC CHECK")
add_text(s, 0.55, 0.55, 12.1, 0.55, "仕様との差分を確認し、残りを M3 の検証に絞った", size=28, bold=True)
add_card(s, 0.65, 1.35, 3.85, 3.45, "満たした", "四隅スタート\n橋と狭い入口\n小さい中央コア\n警告・fade の可読性\n古い不要シーン整理", GREEN)
add_card(s, 4.85, 1.35, 3.85, 3.45, "まだ未検証", "2P 同時プレイ\n体当たり・pulse の納得感\n中心キャンプが強すぎないか\n実動画での見やすさ", ORANGE)
add_card(s, 9.05, 1.35, 3.85, 3.45, "意図的な変更", "縮む速度は一旦遅くした。\n理由は、すぐ終わると接触や判断が見えないため。", TEAL)
add_text(s, 0.95, 5.55, 11.6, 0.42, "仕様照合メモ: docs/reports/2026-06-09_shrinking_island_spec_check.md", size=14, color=MUTED)


# 6 quality
s = new_slide(6, "QUALITY GATE")
add_text(s, 0.55, 0.55, 12.2, 0.55, "直した理由を docs / review に残し、同じ誤診を減らした", size=28, bold=True)
add_card(s, 0.65, 1.30, 5.85, 3.55, "発見した罠", "enabled=false でも public method 呼びは止まらない\nCinemachine が edit-mode の Camera transform を上書き\nMaterialPropertyBlock は reload 後に stale になり得る\nrounded mesh は低角度で side face が崩れる", ORANGE)
add_card(s, 6.85, 1.30, 5.85, 3.55, "今回の確認", "Console: 0\nEditMode: 9 / 9 pass\nRejected mesh refs: 0\nPresentation proxies: inactive by default\nPlay check: P1 grounded", TEAL)
add_text(s, 0.95, 5.55, 11.0, 0.48, "発表での言い方: 失敗を隠さず、同じミスを繰り返さない仕組みに変えた。", size=18, bold=True, color=INK)


# 7 next
s = new_slide(7, "NEXT")
add_text(s, 0.55, 0.55, 12.1, 0.55, "次は M3: 選んだスキルを、縮む島で使う 1 ラウンドを見せる", size=28, bold=True)
add_card(s, 0.65, 1.50, 3.8, 2.2, "1  Draft UI", "ラウンド前に\nスキルを選ぶ", TEAL)
add_card(s, 4.85, 1.50, 3.8, 2.2, "2  Skill Use", "dash / jump / pulse を\n縮む島で使う", ORANGE)
add_card(s, 9.05, 1.50, 3.8, 2.2, "3  2P Playtest", "橋・中央でぶつかり\nKO まで確認する", GREEN)
add_card(s, 0.65, 5.05, 11.95, 0.85, "次の報告目標", "1ラウンドの流れを動画で見せ、スキルが縮む島の事故を増やすかを検証する。", ORANGE)


# 8 thanks
s = new_slide(8)
add_text(s, 0.70, 1.00, 11.8, 0.72, "ありがとうございました", size=44, bold=True)
add_text(s, 0.78, 2.15, 10.8, 0.38, "次回: dash / jump / pulse と 2P Playtest", size=22, color=MUTED)
add_image(s, SHOT, 0.78, 3.05, 5.9, 3.15)


prs.save(OUT)
print(OUT)
